"""Fichas técnicas fictícias no layout em blocos, com documentos de dígito verificador válido.

A terceira ficha tem erros de propósito: autoria que soma 95%, CPF com dígito errado, ISRC vazio,
autor repetido e um rótulo que ninguém conhece (este só gera aviso).
"""
from __future__ import annotations

import random
from datetime import date, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

NOMES = ["Alice", "Benício", "Cecília", "Danilo", "Estela", "Fernando", "Giovana", "Hugo", "Ingrid", "Joaquim",
         "Lívia", "Marcelo", "Natália", "Otávio", "Paloma", "Raul", "Sabrina", "Tomás", "Úrsula", "Valter"]
SOBRENOMES = ["Antunes", "Bezerra", "Coelho", "Damasceno", "Evangelista", "Figueira", "Galvão", "Leite",
              "Monteiro", "Neves", "Paiva", "Ribeiro", "Sampaio", "Teles", "Vasconcelos"]
EDITORAS = ["Editora Pedra Azul", "Edições Lagoa", "Música Serra Verde", "Editora Cais"]
INSTRUMENTOS = ["Violão", "Guitarra", "Baixo", "Bateria", "Teclado", "Percussão", "Voz", "Sanfona", "Cavaquinho"]
TITULOS = ["Rio de Dentro", "Lua de Agosto", "Pé na Areia", "Cais do Porto", "Vento Sul", "Rede na Varanda",
           "Trem das Seis", "Chão Batido", "Fim de Feira", "Céu de Maio", "Barco Lento", "Folha Seca"]
GRAVADORAS = ["Selo Pedra Azul", "Gravadora Lagoa", "Estúdio Cais"]


def cpf(rng: random.Random) -> str:
    base = [rng.randint(0, 9) for _ in range(9)]
    while len(set(base)) == 1:
        base = [rng.randint(0, 9) for _ in range(9)]
    for tamanho in (9, 10):
        soma = sum(d * (tamanho + 1 - i) for i, d in enumerate(base[:tamanho]))
        base.append(soma * 10 % 11 % 10)
    return "".join(map(str, base))


def cnpj(rng: random.Random) -> str:
    base = [rng.randint(0, 9) for _ in range(8)] + [0, 0, 0, 1]
    for pesos in ([5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2], [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]):
        resto = sum(d * p for d, p in zip(base, pesos)) % 11
        base.append(0 if resto < 2 else 11 - resto)
    return "".join(map(str, base))


def _cpf_formatado(numero: str) -> str:
    return f"{numero[:3]}.{numero[3:6]}.{numero[6:9]}-{numero[9:]}"


def _pessoas(rng: random.Random, quantas: int) -> list[tuple[str, str]]:
    nomes = rng.sample([f"{n} {s}" for n in NOMES for s in SOBRENOMES], quantas)
    return [(n, cpf(rng)) for n in nomes]


def gerar(destino: Path | str, semente: int = 3) -> list[Path]:
    destino = Path(destino)
    destino.mkdir(parents=True, exist_ok=True)
    rng = random.Random(semente)
    editoras = [(e, cnpj(rng)) for e in EDITORAS]
    titulos = iter(rng.sample(TITULOS, len(TITULOS)))
    sequencia = iter(range(1, 1000))
    arquivos = []
    for numero, (faixas, com_erros) in enumerate([(5, False), (3, False), (4, True)], start=1):
        caminho = destino / f"ficha_{numero:02d}.xlsx"
        _ficha(caminho, rng, GRAVADORAS[numero - 1], faixas, titulos, sequencia, editoras, com_erros)
        arquivos.append(caminho)
    return arquivos


def _ficha(caminho, rng, gravadora, quantas, titulos, sequencia, editoras, com_erros) -> None:
    livro = Workbook()
    aba = livro.active
    aba.title = "Ficha"
    negrito = Font(bold=True)

    def linha(*valores, destaque=False):
        aba.append(list(valores))
        if destaque:
            aba.cell(row=aba.max_row, column=1).font = negrito

    linha("Ficha técnica de gravação", destaque=True)
    linha("Gravadora:", gravadora)
    linha("Data de envio:", date(2026, 3, 2))
    for n in range(1, quantas + 1):
        linha()
        linha("Faixa", n, destaque=True)
        linha("Título", next(titulos))
        isrc = f"BR-{gravadora[:3].upper()}-26-{next(sequencia):05d}"
        linha("ISRC", None if com_erros and n == 3 else isrc)
        segundos = rng.randint(150, 330)
        duracao = time(0, segundos // 60, segundos % 60) if n % 2 else f"{segundos // 60:02d}:{segundos % 60:02d}"
        linha("Duração", duracao)
        linha("Lançamento", date(2026, 1, 10) + timedelta(days=rng.randint(0, 40)))
        if com_erros and n == 4:
            linha("Observação", "faixa bônus")
        autores = _pessoas(rng, rng.choice([1, 2, 3]))
        partes = {1: [100], 2: [50, 50], 3: [40, 30, 30]}[len(autores)]
        if com_erros and n == 1:
            autores, partes = _pessoas(rng, 2), [50, 45]
        if com_erros and n == 2:
            nome, numero = autores[0]
            autores[0] = (nome, numero[:-1] + str((int(numero[-1]) + 1) % 10))
        if com_erros and n == 4:
            autores, partes = [autores[0], autores[0]], [50, 50]
        linha("Autores", destaque=True)
        linha("Nome", "CPF", "Função", "%")
        for (nome, numero), parte in zip(autores, partes):
            percentual = f"{parte:.2f}".replace(".", ",") if rng.random() < 0.3 else parte
            linha(nome, _cpf_formatado(numero), rng.choice(["Compositor", "Letrista", "Autor"]), percentual)
        if rng.random() < 0.7:
            linha("Editoras", destaque=True)
            linha("Nome", "CNPJ", "%")
            escolhidas = rng.sample(editoras, rng.choice([1, 2]))
            for (nome, numero), parte in zip(escolhidas, [100] if len(escolhidas) == 1 else [60, 40]):
                linha(nome, numero, parte)
        linha("Músicos", destaque=True)
        linha("Nome", "CPF", "Instrumento")
        for nome, numero in _pessoas(rng, rng.randint(2, 5)):
            linha(nome, numero, rng.choice(INSTRUMENTOS))
    aba.column_dimensions["A"].width = 28
    aba.column_dimensions["B"].width = 22
    livro.save(caminho)
