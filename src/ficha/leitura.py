"""Leitura da ficha em Excel, que não é uma tabela: é uma sequência de blocos, um por faixa.

Cada bloco começa na linha "Faixa", tem campos no formato rótulo/valor (Título, ISRC, Duração,
Lançamento) e três seções em tabela (Autores, Editoras, Músicos). A leitura anda linha a linha, como
uma máquina de estados, e guarda a célula de origem de cada valor para o erro apontar o lugar certo.
"""
from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from .modelo import Faixa, Ficha, Pessoa, Problema

CAMPOS_FAIXA = {"TITULO": "titulo", "ISRC": "isrc", "DURACAO": "duracao_seg", "LANCAMENTO": "lancamento"}
SECOES = {"AUTORES": "autores", "EDITORAS": "editoras", "MUSICOS": "musicos"}


def rotulo(valor) -> str:
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", texto).strip().rstrip(":").upper()


def ler(caminho: Path | str) -> tuple[Ficha, list[Problema]]:
    aba = load_workbook(caminho, data_only=True).active
    ficha, problemas = Ficha(), []
    faixa: Faixa | None = None
    secao: str | None = None
    for linha in aba.iter_rows():
        celulas = [c for c in linha if c.value is not None and str(c.value).strip() != ""]
        if not celulas:
            secao = None  # linha em branco fecha a seção
            continue
        a, ref = linha[0].value, linha[0].coordinate
        valor = linha[1].value if len(linha) > 1 else None
        chave = rotulo(a)

        if chave in ("FICHA TECNICA DE GRAVACAO", "NOME"):
            continue
        if chave == "GRAVADORA":
            ficha.gravadora = _texto(valor)
        elif chave == "DATA DE ENVIO":
            ficha.data_envio = _data(valor, f"B{linha[0].row}", problemas)
        elif chave == "FAIXA":
            faixa = Faixa(numero=_inteiro(valor, f"B{linha[0].row}", problemas), celula=ref)
            ficha.faixas.append(faixa)
            secao = None
        elif faixa is None:
            problemas.append(Problema(ref, f"conteúdo antes da primeira faixa: {a!r}", grave=False))
        elif chave in CAMPOS_FAIXA:
            celula_valor = f"B{linha[0].row}"
            faixa.celulas[CAMPOS_FAIXA[chave]] = celula_valor
            if chave == "DURACAO":
                faixa.duracao_seg = _duracao(valor, celula_valor, problemas)
            elif chave == "LANCAMENTO":
                faixa.lancamento = _data(valor, celula_valor, problemas)
            else:
                setattr(faixa, CAMPOS_FAIXA[chave], _texto(valor))
        elif chave in SECOES:
            secao = SECOES[chave]
            faixa.celulas[secao] = ref
        elif secao:
            getattr(faixa, secao).append(_pessoa(secao, [c.value for c in linha[:4]], ref, problemas))
        else:
            problemas.append(Problema(ref, f"rótulo desconhecido: {a!r}", grave=False))
    return ficha, problemas


def _pessoa(secao: str, valores: list, ref: str, problemas: list[Problema]) -> Pessoa:
    valores += [None] * (4 - len(valores))
    nome, documento, terceiro, quarto = valores
    p = Pessoa(nome=_texto(nome) or "", documento=re.sub(r"\D", "", str(documento or "")), celula=ref)
    if secao == "autores":
        p.funcao, p.percentual = rotulo(terceiro) or None, _percentual(quarto, ref, problemas)
    elif secao == "editoras":
        p.percentual = _percentual(terceiro, ref, problemas)
    else:
        p.instrumento = rotulo(terceiro) or None
    return p


def _texto(valor) -> str | None:
    texto = re.sub(r"\s+", " ", str(valor)).strip() if valor is not None else ""
    return texto or None


def _inteiro(valor, ref: str, problemas: list[Problema]) -> int:
    try:
        return int(valor)
    except (TypeError, ValueError):
        problemas.append(Problema(ref, f"número da faixa inválido: {valor!r}"))
        return 0


def _data(valor, ref: str, problemas: list[Problema]) -> date | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    try:
        return datetime.strptime(str(valor).strip(), "%d/%m/%Y").date()
    except ValueError:
        problemas.append(Problema(ref, f"data inválida: {valor!r} (use dd/mm/aaaa)"))
        return None


def _duracao(valor, ref: str, problemas: list[Problema]) -> int | None:
    """Aceita a célula formatada como hora do Excel ou texto mm:ss / h:mm:ss."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, time):
        return valor.hour * 3600 + valor.minute * 60 + valor.second
    partes = str(valor).strip().split(":")
    if 2 <= len(partes) <= 3 and all(p.isdigit() for p in partes):
        numeros = [int(p) for p in partes]
        horas, minutos, segundos = ([0] + numeros)[-3:]
        if minutos < 60 and segundos < 60:
            return horas * 3600 + minutos * 60 + segundos
    problemas.append(Problema(ref, f"duração inválida: {valor!r} (use mm:ss)"))
    return None


def _percentual(valor, ref: str, problemas: list[Problema]) -> Decimal | None:
    if valor is None or valor == "":
        return None
    texto = str(valor).strip().rstrip("%").strip()
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        return Decimal(texto).quantize(Decimal("0.01"))
    except InvalidOperation:
        problemas.append(Problema(ref, f"percentual inválido: {valor!r}"))
        return None
